"""
Read an HTM (Hazard Traceability Matrix) xlsx using the schema map declared
in project.yml ``dhfs[].evidence.hazard_traceability_matrix[]`` and produce
a normalized HtmState consumed by the drift-rule registry.

Project-agnostic by construction. Schema-driven the same way `dtm_reader.py`
is — every column, extractor, and TBD sentinel comes from `HtmSchema`
(resolved by `lib.config`). This module does not know the prefix of any
particular project's hazard IDs.

Defensive contract:

* If a project has no HTM schema declared (the common case during early
  development — HTM Confluence page is a placeholder), audit.py passes
  ``None`` to every rule and the hazard-edge rules degrade gracefully.
* If a schema is declared but the xlsx file is missing, ``read`` raises
  ``FileNotFoundError`` so the audit runner can surface the misconfig
  rather than silently audit against an empty matrix.
"""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Pattern

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from .config import HtmSchema


@dataclass(frozen=True)
class HtmRow:
    """One normalized HTM data row.

    Conceptually each row links a Hazard to one or more Design Inputs and
    optionally to a mitigation (control measure) and verification evidence.
    """
    row_index: int
    cells: Dict[str, Optional[str]]
    extracted: Dict[str, Optional[str]]
    is_blank: bool


@dataclass(frozen=True)
class HtmState:
    schema: HtmSchema
    rows: List[HtmRow]
    sheet_name: str
    source_path: str

    # Convenience indices
    by_hazard: Dict[str, List[HtmRow]] = field(default_factory=dict)
    by_design_input: Dict[str, List[HtmRow]] = field(default_factory=dict)


def _is_blank(value: Optional[str], schema: HtmSchema) -> bool:
    if value is None:
        return True
    s = str(value).strip()
    if s == "":
        return True
    if s == schema.tbd_value:
        return True
    return False


def _cell_text(value) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def _normalize_id(raw: str) -> str:
    return re.sub(r"[\s-]+", "", raw.strip())


def _compile_extractors(schema: HtmSchema) -> Dict[str, Pattern[str]]:
    return {name: re.compile(pat) for name, pat in (schema.extractors or {}).items()}


def _extract_first(text: Optional[str], pattern: Pattern[str]) -> Optional[str]:
    if text is None:
        return None
    m = pattern.search(text)
    if not m:
        return None
    raw = m.group(1) if m.lastindex else m.group(0)
    return _normalize_id(raw)


def _extract_all(text: Optional[str], pattern: Pattern[str]) -> List[str]:
    if text is None:
        return []
    return [
        _normalize_id(m.group(1) if m.lastindex else m.group(0))
        for m in pattern.finditer(text)
    ]


def read(schema: HtmSchema, project_root: Path) -> HtmState:
    """Read the HTM xlsx into an HtmState. Path resolved relative to project_root.

    Raises FileNotFoundError when the schema points at a missing file —
    that is a misconfig the runner should surface rather than mask.
    """
    xlsx_path = (project_root / schema.xlsx_path).resolve()
    if not xlsx_path.exists():
        raise FileNotFoundError(f"HTM xlsx not found: {xlsx_path}")
    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    if schema.sheet not in wb.sheetnames:
        raise KeyError(
            f"HTM sheet {schema.sheet!r} not in {xlsx_path.name}; sheets: {wb.sheetnames}"
        )
    ws = wb[schema.sheet]

    column_indices = {name: column_index_from_string(letter)
                      for name, letter in schema.columns.items()}
    extractor_patterns = _compile_extractors(schema)

    rows: List[HtmRow] = []
    by_haz: Dict[str, List[HtmRow]] = {}
    by_di: Dict[str, List[HtmRow]] = {}

    for excel_row_num, raw_row in enumerate(
        ws.iter_rows(min_row=schema.data_start_row, values_only=True),
        start=schema.data_start_row,
    ):
        cells: Dict[str, Optional[str]] = {}
        for col_name, col_idx in column_indices.items():
            value = raw_row[col_idx - 1] if col_idx - 1 < len(raw_row) else None
            cells[col_name] = _cell_text(value)

        is_blank = all(_is_blank(cells[name], schema) for name in column_indices)
        if is_blank:
            continue

        extracted: Dict[str, Optional[str]] = {}
        for name, pat in extractor_patterns.items():
            source_col = name.replace("_id", "")
            extracted[name] = _extract_first(cells.get(source_col), pat)

        row = HtmRow(
            row_index=excel_row_num,
            cells=cells,
            extracted=extracted,
            is_blank=False,
        )
        rows.append(row)

        if hz := extracted.get("hazard_id"):
            by_haz.setdefault(hz, []).append(row)
        di_pat = extractor_patterns.get("design_input_id")
        if di_pat is not None:
            for di in _extract_all(cells.get("design_input"), di_pat):
                by_di.setdefault(di, []).append(row)

    wb.close()

    return HtmState(
        schema=schema,
        rows=rows,
        sheet_name=schema.sheet,
        source_path=str(xlsx_path),
        by_hazard=by_haz,
        by_design_input=by_di,
    )
